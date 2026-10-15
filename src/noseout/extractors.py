"""파일에서 색인할 내용을 꺼낸다: 문서 텍스트 청크, PDF 페이지(텍스트 또는 렌더링한 그림), 이미지 정보.
원본 Extractors.swift, PDF+Portable.swift, Images+Portable.swift의 동작을 따른다. 파일은 읽기만 한다."""

from __future__ import annotations

import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image, ImageOps

from . import rules

try:
    import pillow_heif

    pillow_heif.register_heif_opener()
except ImportError:  # HEIC 없이도 나머지는 동작
    pass

CHUNK_SIZE = 1800
CHUNK_OVERLAP = 200
MAX_CHUNKS = 5000
MAX_PAGES = 5000
MIN_PAGE_TEXT = 200       # 이보다 글자가 적은 페이지는 그림으로 렌더링해 이미지 벡터를 만든다
PAGE_TEXT_LIMIT = 4000    # 페이지 텍스트 임베딩에 쓰는 앞부분
RENDER_PIXELS = 1400      # 렌더링한 페이지의 긴 변
# 텍스트 파일은 앞부분만 읽는다. MAX_CHUNKS × CHUNK_SIZE(900만 자)를 덮는 크기라 결과는 같고, 수 GB짜리 덤프가
# 메모리를 다 쓰지 않는다.
MAX_TEXT_BYTES = 32 * 1024 * 1024


class ExtractError(Exception):
    pass


class LeftOut(ExtractError):
    """읽을 수는 있지만 색인하지 않기로 한 파일(코드 색인: 바이너리, 생성·압축된 코드, 개인 키). failed가 아니라 skipped."""


# 문서


def chunks(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP, limit: int = MAX_CHUNKS) -> list[str]:
    """~size자 조각, 조금씩 겹치고 공백에서 끊는다 (원본 Extractors.chunks와 같은 규칙)."""
    clean = re.sub(r"[ \t ]+", " ", text)
    clean = re.sub(r"\n\s*\n\s*\n+", "\n\n", clean).strip()
    out: list[str] = []
    start = 0
    n = len(clean)
    while start < n and len(out) < limit:
        end = min(start + size, n)
        if end < n:
            space = max((i for i in range(start, end) if clean[i].isspace()), default=-1)
            if space - start > size // 2:
                end = space
        piece = clean[start:end].strip()
        if piece:
            out.append(piece)
        if end >= n:
            break
        nxt = max(end - overlap, start)
        if nxt <= start:
            nxt = end
        while nxt < end and not clean[nxt].isspace():
            nxt += 1
        start = nxt if nxt < end else end
    return out


def read_text(path: str) -> str:
    """txt/md/html/docx의 본문."""
    ext = rules.extension(path)
    if ext == "docx":
        return _docx_text(path)
    with open(path, "rb") as f:
        raw = f.read(MAX_TEXT_BYTES)
    text = _decode(raw, cut=len(raw) == MAX_TEXT_BYTES)
    if ext in ("html", "htm"):
        text = _html_text(text)
    return text


def _decode(raw: bytes, cut: bool = False) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw[: len(raw) // 2 * 2].decode("utf-16", errors="replace")
    # 앞부분만 읽었으면 끝에서 UTF-8 글자가 잘렸을 수 있다: 최대 3바이트까지 떼고 다시 해 본다
    for drop in range(4 if cut else 1):
        try:
            return raw[: len(raw) - drop].decode("utf-8")
        except UnicodeDecodeError:
            continue
    return raw.decode("cp949", errors="replace")  # 한국어 Windows의 옛 메모장 파일


class _HTMLText(HTMLParser):
    """본문 글자만 모은다. script·style·head 안은 건너뛴다. 앞에서부터 한 번 훑으므로 닫는 태그가 없는 큰 파일도
    선형 시간이다 (전에 쓰던 정규식 `<script>.*?</script>`는 닫는 태그가 없으면 이차 시간이 걸렸다)."""

    SKIP = {"script", "style", "head"}
    BREAK = {"p", "div", "li", "br", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag in self.BREAK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skipping = max(0, self.skipping - 1)
        elif tag in self.BREAK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skipping:
            self.parts.append(data)


def _html_text(source: str) -> str:
    parser = _HTMLText()
    parser.feed(source)
    parser.close()
    return " ".join(parser.parts)


def _docx_text(path: str) -> str:
    import docx  # python-docx

    try:
        document = docx.Document(path)
    except Exception as error:  # 암호가 걸렸거나 깨진 파일
        raise ExtractError(f"unreadable docx: {error}") from None
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


# PDF


@dataclass
class Page:
    number: int              # 1부터
    text: str
    image_path: str | None   # 글자가 적어 렌더링했을 때


def pdf_pages(path: str, render_to: Path, start: int = 0, stop: int = MAX_PAGES) -> tuple[list[Page], int, str | None]:
    """페이지 `start`..`stop`(0부터, stop 미포함)의 (페이지들, 전체 페이지 수, 문서 제목). 긴 PDF는 나눠 읽는다."""
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as error:
        raise ExtractError(f"unreadable pdf: {error}") from None
    try:
        total = len(pdf)
        title = (pdf.get_metadata_dict().get("Title") or "").strip() or None
        pages = [_read_page(pdf, index, render_to) for index in range(start, min(total, stop, MAX_PAGES))]
        return pages, total, title
    except pdfium.PdfiumError as error:
        # 문서는 열렸지만 페이지가 깨진 경우: 이 파일만 실패로 두고 나머지 파일은 계속 (sync 전체를 멈추지 않게)
        raise ExtractError(f"unreadable pdf page: {error}") from None
    finally:
        pdf.close()


def _read_page(pdf, index: int, render_to: Path) -> Page:
    page = pdf[index]
    try:
        textpage = page.get_textpage()
        text = textpage.get_text_bounded().replace("\r\n", "\n").strip()
        textpage.close()
        image_path = None
        if len(text) < MIN_PAGE_TEXT:
            width, height = page.get_size()
            bitmap = page.render(scale=RENDER_PIXELS / max(width, height))
            render_to.mkdir(parents=True, exist_ok=True)
            image_path = str(render_to / f"page-{index + 1}.png")
            bitmap.to_pil().convert("RGB").save(image_path)
            bitmap.close()
        return Page(index + 1, text, image_path)
    finally:
        page.close()


# 이미지


@dataclass
class ImageInfo:
    width: int
    height: int
    taken: str | None


def image_info(path: str) -> ImageInfo:
    """보이는 방향(EXIF 회전 적용) 크기와 촬영 시각. 너무 작은 그림(아이콘)은 ExtractError."""
    try:
        with Image.open(path) as image:
            exif = image.getexif()
            taken = exif.get_ifd(0x8769).get(0x9003) or exif.get(0x0132)
            oriented = ImageOps.exif_transpose(image)
            width, height = oriented.size
    except Exception as error:
        raise ExtractError(f"unreadable image: {error}") from None
    if max(width, height) < rules.MIN_IMAGE_SIDE:
        raise ExtractError(f"small image {width}x{height}")
    return ImageInfo(width, height, str(taken) if taken else None)
